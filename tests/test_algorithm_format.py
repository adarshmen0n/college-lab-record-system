"""
Comprehensive tests for strict Step-by-Step Algorithm format:
- Dynamic step numbering (from 1 to 25+ steps)
- Prefix normalization: Step {step_number}: {instruction}
- Input normalization (plain text, numbered, bullets, existing Step prefixes, double prefixes)
- DOCX typography (Times New Roman 12pt Regular, 0.4" indent, 1.15 line spacing)
- Heading ALGORITHM: (Times New Roman 14pt Bold)
- Schema, PageEngine, DocxGenerator, and API synchronization
"""
import os
import pytest
import docx
from docx.shared import Inches, Pt
from fastapi.testclient import TestClient

from backend.main import app
from backend.models.schemas import ExperimentData, TemplateConfig
from backend.sanitizer import normalize_algorithm_step, format_algorithm_steps
from backend.layout.page_engine import PageEngine
from backend.generator.docx_generator import DocxGenerator
from backend.validator.page_checker import PageChecker

client = TestClient(app)

@pytest.fixture
def base_template():
    return TemplateConfig()

def test_normalization_plain_lines():
    raw = "Start\nRead input\nCalculate result"
    steps = format_algorithm_steps(raw)
    assert len(steps) == 3
    assert steps[0] == "Step 1: Start."
    assert steps[1] == "Step 2: Read input."
    assert steps[2] == "Step 3: Calculate result."

def test_normalization_numbered_lines():
    raw = "1. Start\n2. Read input\n3. Calculate result"
    steps = format_algorithm_steps(raw)
    assert len(steps) == 3
    assert steps[0] == "Step 1: Start."
    assert steps[1] == "Step 2: Read input."
    assert steps[2] == "Step 3: Calculate result."
    # Prevent duplicate numbering like "Step 1: 1. Start"
    assert "1." not in steps[0].split("Step 1:")[1]

def test_normalization_bulleted_lines():
    raw = "- Start\n* Read input\n• Calculate result"
    steps = format_algorithm_steps(raw)
    assert len(steps) == 3
    assert steps[0] == "Step 1: Start."
    assert steps[1] == "Step 2: Read input."
    assert steps[2] == "Step 3: Calculate result."

def test_normalization_already_step_prefixed():
    raw = "Step 1: Start\nStep 2: Read input\nStep 3: Calculate result"
    steps = format_algorithm_steps(raw)
    assert len(steps) == 3
    assert steps[0] == "Step 1: Start."
    assert steps[1] == "Step 2: Read input."
    assert steps[2] == "Step 3: Calculate result."
    # Prevent duplicate Step prefix like "Step 1: Step 1: Start"
    assert "Step 1: Step 1:" not in steps[0]

def test_normalization_double_prefix():
    raw = "Step 1: 1. Start\nStep 2 - 2) Read input\nStep 3. (3) Calculate result"
    steps = format_algorithm_steps(raw)
    assert len(steps) == 3
    assert steps[0] == "Step 1: Start."
    assert steps[1] == "Step 2: Read input."
    assert steps[2] == "Step 3: Calculate result."

def test_punctuation_preservation():
    raw = "Does condition hold?\nStop immediately!\nReturn total"
    steps = format_algorithm_steps(raw)
    assert steps[0] == "Step 1: Does condition hold?"
    assert steps[1] == "Step 2: Stop immediately!"
    assert steps[2] == "Step 3: Return total."

def test_dynamic_step_numbering_requirement_13():
    """Exact sample from Requirement 13 in prompt: 6 steps."""
    raw = """Create an empty ordered dictionary.
Display the menu with options Insert, Display, Search and Exit.
Insert a key-value pair into the dictionary.
Display all the key-value pairs.
Search for a particular key in the dictionary.
Exit the program."""
    steps = format_algorithm_steps(raw)
    assert len(steps) == 6
    assert steps[0] == "Step 1: Create an empty ordered dictionary."
    assert steps[1] == "Step 2: Display the menu with options Insert, Display, Search and Exit."
    assert steps[2] == "Step 3: Insert a key-value pair into the dictionary."
    assert steps[3] == "Step 4: Display all the key-value pairs."
    assert steps[4] == "Step 5: Search for a particular key in the dictionary."
    assert steps[5] == "Step 6: Exit the program."

def test_dynamic_large_step_count():
    """Verify system dynamically scales to 25+ steps without hardcoded cap."""
    raw = "\n".join([f"Process item number {i} in sequence" for i in range(1, 26)])
    steps = format_algorithm_steps(raw)
    assert len(steps) == 25
    assert steps[0] == "Step 1: Process item number 1 in sequence."
    assert steps[9] == "Step 10: Process item number 10 in sequence."
    assert steps[24] == "Step 25: Process item number 25 in sequence."

def test_schema_automatic_normalization():
    exp = ExperimentData(
        experiment_number="3.A",
        title="ORDERED DICTIONARY",
        aim="To implement an ordered dictionary in Python.",
        algorithm="1. Initialize dictionary\n2. Add entries\n3. Display entries\n4. Exit",
        coding="from collections import OrderedDict\nd = OrderedDict()\nprint(d)",
        output="OrderedDict()",
        result="Ordered dictionary demonstrated successfully."
    )
    assert "Step 1: Initialize dictionary." in exp.algorithm
    assert "Step 2: Add entries." in exp.algorithm
    assert "Step 3: Display entries." in exp.algorithm
    assert "Step 4: Exit." in exp.algorithm

def test_page_engine_algorithm_steps(base_template):
    exp = ExperimentData(
        experiment_number="3.A",
        title="ORDERED DICTIONARY",
        aim="To implement an ordered dictionary in Python.",
        algorithm="1. Step one\n2. Step two\n3. Step three",
        coding="print('test')",
        output="test",
        result="Success"
    )
    plan = PageEngine.plan_experiment(exp, base_template)
    p1 = plan.pages[0]
    assert len(p1.algorithm_steps) == 3
    assert p1.algorithm_steps[0] == "Step 1: Step one."
    assert p1.algorithm_steps[1] == "Step 2: Step two."
    assert p1.algorithm_steps[2] == "Step 3: Step three."

def test_docx_algorithm_typography(tmp_path, base_template):
    """Verify DOCX paragraph indentation, line spacing, font size, and typography."""
    raw_algo = """Create an empty ordered dictionary.
Display the menu with options Insert, Display, Search and Exit.
Insert a key-value pair into the dictionary.
Display all the key-value pairs.
Search for a particular key in the dictionary.
Exit the program."""

    exp = ExperimentData(
        experiment_number="3.A",
        title="ORDERED DICTIONARY OPERATIONS",
        aim="To implement and test operations on an ordered dictionary.",
        algorithm=raw_algo,
        coding="from collections import OrderedDict\n\ndef main():\n    d = OrderedDict()\n    print(d)",
        output="OrderedDict()",
        result="Ordered dictionary operations executed successfully."
    )

    out_file = str(tmp_path / "test_algo_typography.docx")
    DocxGenerator.generate_new_record(exp, base_template, out_file)
    assert os.path.exists(out_file)

    doc = docx.Document(out_file)
    
    # Locate ALGORITHM heading and following step paragraphs
    algo_heading_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip().upper() == "ALGORITHM:":
            algo_heading_idx = idx
            break

    assert algo_heading_idx is not None, "ALGORITHM: heading not found"
    heading_p = doc.paragraphs[algo_heading_idx]
    assert heading_p.runs[0].font.name == "Times New Roman"
    assert round(heading_p.runs[0].font.size.pt) == 14
    assert heading_p.runs[0].font.bold is True

    # Check the 6 step paragraphs
    step_paragraphs = doc.paragraphs[algo_heading_idx + 1: algo_heading_idx + 7]
    assert len(step_paragraphs) == 6

    for i, sp in enumerate(step_paragraphs, 1):
        assert sp.text.startswith(f"Step {i}:")
        assert sp.runs[0].font.name == "Times New Roman"
        assert round(sp.runs[0].font.size.pt) == 12
        assert sp.runs[0].font.bold in (False, None)  # Regular font
        # Indent: 0.4 inches (365760 EMUs)
        assert sp.paragraph_format.left_indent == Inches(0.4)
        # Line spacing 1.15
        assert sp.paragraph_format.line_spacing == 1.15
        # Clean paragraph spacing
        assert sp.paragraph_format.space_before == Pt(2)
        assert sp.paragraph_format.space_after == Pt(3)

    # Validate overall document structure
    val = PageChecker.validate_document(out_file, exp)
    assert val.is_valid is True
    assert val.checks["no_instruction_leakage"] is True
    assert val.checks["has_header_table"] is True
    assert val.checks["has_evaluation_table"] is True
