"""
Tests for Multi-Experiment Workspace, Page-by-Page Preview, and Compact Header Table.
"""
import os
import pytest
from fastapi.testclient import TestClient
import docx
from docx.shared import Pt

from backend.main import app
from backend.models.schemas import ExperimentData, TemplateConfig
from backend.generator.docx_generator import DocxGenerator
from backend.generator.tables import TableManager

client = TestClient(app)

@pytest.fixture
def sample_exp_1():
    return {
        "experiment_number": "3.A",
        "title": "STUDY OF BASIC PYTHON SYNTAX AND OPERATIONS",
        "subtitle": "PYTHON LAB",
        "date": "23-09-2026",
        "aim": "To write and execute basic Python programs demonstrating variables, operators, and data types.",
        "procedure_heading": "ALGORITHM",
        "algorithm": "1. Start the program.\n2. Initialize variables with sample values.\n3. Perform arithmetic operations.\n4. Display the results.\n5. Stop the program.",
        "code_heading": "CODING",
        "coding": "a = 10\nb = 20\nprint('Sum:', a + b)\nprint('Product:', a * b)",
        "output": "Sum: 30\nProduct: 200",
        "result": "Thus, the basic Python programs were written and executed successfully.",
        "student_name": "ADARSH MENON",
        "register_number": "714025247005"
    }

@pytest.fixture
def sample_exp_2():
    return {
        "experiment_number": "3.B",
        "title": "IMPLEMENTATION OF CONTROL FLOW STATEMENTS",
        "subtitle": "PYTHON LAB",
        "date": "24-09-2026",
        "aim": "To implement control flow statements including if-else and loops in Python.",
        "procedure_heading": "ALGORITHM",
        "algorithm": "1. Start the program.\n2. Read input value.\n3. Check condition using if-else.\n4. Print conditional output.\n5. Stop.",
        "code_heading": "CODING",
        "coding": "x = 15\nif x % 2 == 0:\n    print('Even')\nelse:\n    print('Odd')",
        "output": "Odd",
        "result": "Thus, the control flow statements were implemented and verified.",
        "student_name": "ADARSH MENON",
        "register_number": "714025247005"
    }

def test_compact_header_table_structure(tmp_path):
    doc = docx.Document()
    cfg = TemplateConfig()
    exp = ExperimentData(
        experiment_number="3.A",
        title="STUDY OF BASIC PYTHON SYNTAX AND OPERATIONS",
        subtitle="",
        date="23-09-2026",
        aim="Test aim",
        algorithm="1. Test",
        coding="print('test')",
        output="test",
        result="Test passed",
        student_name="ADARSH MENON",
        register_number="714025247005"
    )

    tbl = TableManager.create_header_table(doc, exp, cfg)
    assert tbl is not None
    # 2 rows, 2 columns in underlying grid
    assert len(tbl.rows) == 2
    assert len(tbl.columns) == 2

    # Left cells
    assert "EX NO:3.A" in tbl.cell(0, 0).text
    assert "DATE:23-09-2026" in tbl.cell(1, 0).text

    # Merged right cell check
    right_cell_0 = tbl.cell(0, 1)
    right_cell_1 = tbl.cell(1, 1)
    # Merged cells in python-docx reference the same underlying XML element or share text
    assert "STUDY OF BASIC PYTHON SYNTAX AND OPERATIONS" in right_cell_0.text
    assert right_cell_0.text == right_cell_1.text

    # Title run font size must be 14pt
    p_title = right_cell_0.paragraphs[0]
    assert len(p_title.runs) > 0
    assert p_title.runs[0].font.name == "Times New Roman"
    assert p_title.runs[0].font.size == Pt(14)
    assert p_title.runs[0].font.bold is True

def test_typography_in_generated_record(tmp_path):
    out_file = str(tmp_path / "test_typo.docx")
    cfg = TemplateConfig()
    exp = ExperimentData(
        experiment_number="1.A",
        title="SAMPLE EXPERIMENT",
        date="2026-09-23",
        aim="To verify typography.",
        algorithm="Step 1\nStep 2",
        coding="x = 10\nprint(x)",
        output="10",
        result="Success",
        student_name="ADARSH MENON",
        register_number="714025247005"
    )
    DocxGenerator.generate_new_record(exp, cfg, out_file)
    assert os.path.exists(out_file)

    doc = docx.Document(out_file)

    # Check headings are 14pt Bold Times New Roman
    headings = ["AIM:", "ALGORITHM:", "CODING:", "OUTPUT:", "RESULT:"]
    found_headings = []
    for p in doc.paragraphs:
        p_text = p.text.strip()
        for h in headings:
            if p_text == h:
                found_headings.append(h)
                for run in p.runs:
                    if run.text.strip():
                        assert run.font.name == "Times New Roman"
                        assert run.font.size == Pt(14)
                        assert run.font.bold is True

    assert "AIM:" in found_headings
    assert "RESULT:" in found_headings

def test_workspace_preview_endpoint(sample_exp_1, sample_exp_2):
    payload = {
        "experiments": [sample_exp_1, sample_exp_2],
        "template_id": "python_lab_reference"
    }
    resp = client.post("/api/record/preview-workspace", json=payload)
    assert resp.status_code == 200
    data = resp.json()

    assert data["success"] is True
    assert data["total_experiments"] == 2
    assert data["total_pages"] == 8
    pages = data["pages"]
    assert len(pages) == 8

    # Page 1 (Exp 1 Start - Right)
    assert pages[0]["page_number"] == 1
    assert pages[0]["side"] == "RIGHT"
    assert pages[0]["has_header_table"] is True
    assert pages[0]["header_ex_no"] == "3.A"
    assert pages[0]["aim_heading"] == "AIM:"

    # Page 2 (Exp 1 Output - Left)
    assert pages[1]["page_number"] == 2
    assert pages[1]["side"] == "LEFT"
    assert pages[1]["output_heading"] == "OUTPUT:"

    # Page 3 (Exp 1 Cont - Right)
    assert pages[2]["page_number"] == 3
    assert pages[2]["side"] == "RIGHT"
    assert pages[2]["has_evaluation_table"] is True
    assert pages[2]["result_heading"] == "RESULT:"

    # Page 4 (Exp 1 Blank Back - Left)
    assert pages[3]["page_number"] == 4
    assert pages[3]["side"] == "LEFT"
    assert pages[3]["is_blank"] is True

    # Page 5 (Exp 2 Start - Right)
    assert pages[4]["page_number"] == 5
    assert pages[4]["side"] == "RIGHT"
    assert pages[4]["has_header_table"] is True
    assert pages[4]["header_ex_no"] == "3.B"

def test_workspace_generate_endpoint(sample_exp_1, sample_exp_2):
    payload = {
        "experiments": [sample_exp_1, sample_exp_2],
        "template_id": "python_lab_reference"
    }
    resp = client.post("/api/record/generate-workspace", json=payload)
    assert resp.status_code == 200
    data = resp.json()

    assert data["success"] is True
    assert "Lab_Record_Exps_3_A_to_3_B_" in data["filename"]
    assert data["total_pages_estimated"] == 8
    assert data["is_continuation"] is False

    # Download file and inspect
    dl_resp = client.get(data["download_url"])
    assert dl_resp.status_code == 200
    assert len(dl_resp.content) > 0
