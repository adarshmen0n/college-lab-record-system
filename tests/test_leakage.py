"""
Regression Tests for Prompt, System Instruction, and Developer Content Leakage Prevention.
Guarantees that internal instructions, system prompts, or developer meta-prompts
NEVER enter the generated laboratory record document.
"""
import os
import pytest
import docx

from backend.models.schemas import ExperimentData, TemplateConfig
from backend.generator.docx_generator import DocxGenerator
from backend.validator.page_checker import PageChecker

FORBIDDEN_PHRASES = [
    "MASTER UPDATE PROMPT",
    "FINAL MASTER EXECUTION PROMPT",
    "PROJECT UPDATE — EXPERIMENT HEADER",
    "You are modifying the existing",
    "You are working on",
    "Do not rebuild",
    "Antigravity",
    "system prompt",
    "developer instructions",
    "implementation instructions"
]

def test_clean_experiment_has_no_instruction_leakage(tmp_path):
    """Verifies that normal generated experiments contain zero prompt or developer leakage."""
    out_file = str(tmp_path / "clean_record.docx")
    cfg = TemplateConfig()
    exp = ExperimentData(
        experiment_number="3.A",
        title="Menu-Driven Program Using Ordered Dictionary",
        date="23-09-2026",
        aim="To perform insert, display, and search operations using an ordered dictionary in Python.",
        procedure_heading="ALGORITHM",
        algorithm="1. Create an empty ordered dictionary.\n2. Display menu.\n3. Stop.",
        code_heading="PROGRAM",
        coding="from collections import OrderedDict\n\ndef menu():\n    data = OrderedDict()",
        output="1. Insert\n2. Display\n3. Exit",
        result="The menu-driven program was executed successfully.",
        student_name="ADARSH MENON",
        register_number="714025247005"
    )

    DocxGenerator.generate_new_record(exp, cfg, out_file)
    assert os.path.exists(out_file)

    doc = docx.Document(out_file)
    full_text = " ".join([p.text for p in doc.paragraphs])
    for tbl in doc.tables:
        full_text += " " + " ".join([c.text for r in tbl.rows for c in r.cells])

    # Search for all forbidden phrases
    for phrase in FORBIDDEN_PHRASES:
        assert phrase.upper() not in full_text.upper(), f"Forbidden phrase '{phrase}' found in generated document!"

    # Validate with PageChecker
    report = PageChecker.validate_document(out_file, exp)
    assert report.is_valid is True
    assert report.checks.get("no_instruction_leakage") is True

def test_prompt_sanitization_strips_injected_prompts(tmp_path):
    """
    Verifies that if prompt instructions or developer text accidentally enter
    any experiment fields, DocxGenerator strips them cleanly before rendering.
    """
    out_file = str(tmp_path / "sanitized_record.docx")
    cfg = TemplateConfig()
    
    # Intentionally inject developer instructions into coding and algorithm
    injected_code = (
        "# MASTER UPDATE PROMPT\n"
        "# You are working on the existing General College Laboratory Record Automation System\n"
        "# Do not rebuild the project\n"
        "from collections import OrderedDict\n"
        "data = OrderedDict()"
    )
    
    exp = ExperimentData(
        experiment_number="3.A",
        title="Menu-Driven Program Using Ordered Dictionary",
        date="23-09-2026",
        aim="To perform insert, display, and search operations using an ordered dictionary in Python.",
        procedure_heading="ALGORITHM",
        algorithm="1. Create dictionary.\n2. MASTER UPDATE PROMPT\n3. Exit.",
        code_heading="PROGRAM",
        coding=injected_code,
        output="Success",
        result="Successfully executed.",
        student_name="ADARSH MENON",
        register_number="714025247005"
    )

    DocxGenerator.generate_new_record(exp, cfg, out_file)
    assert os.path.exists(out_file)

    doc = docx.Document(out_file)
    full_text = " ".join([p.text for p in doc.paragraphs])
    for tbl in doc.tables:
        full_text += " " + " ".join([c.text for r in tbl.rows for c in r.cells])

    # Verify that forbidden phrases were stripped
    assert "MASTER UPDATE PROMPT" not in full_text
    assert "You are working on the existing" not in full_text
    assert "Do not rebuild the project" not in full_text

    # Real code content remains preserved
    assert "from collections import OrderedDict" in full_text
    assert "data = OrderedDict()" in full_text

    # Validate with PageChecker
    report = PageChecker.validate_document(out_file, exp)
    assert report.is_valid is True
    assert report.checks.get("no_instruction_leakage") is True

def test_page_checker_detects_and_fails_on_leakage(tmp_path):
    """Verifies that PageChecker detects and fails validation if a document directly contains forbidden prompts."""
    doc_path = str(tmp_path / "leaked_doc.docx")
    doc = docx.Document()
    doc.add_paragraph("MASTER UPDATE PROMPT")
    doc.add_paragraph("You are working on the existing General College Laboratory Record Automation System")
    doc.save(doc_path)

    report = PageChecker.validate_document(doc_path)
    assert report.is_valid is False
    assert report.checks.get("no_instruction_leakage") is False
    assert any("leakage detected" in err.lower() for err in report.errors)

def test_user_reported_prompt_leakage_fully_stripped(tmp_path):
    """
    Specifically tests the exact user-reported prompt text from prompt #10:
    '# PROJECT UPDATE — EXPERIMENT HEADER + TYPOGRAPHY CORRECTION
    You are modifying the existing General College Laboratory Record Automation System.
    Do NOT rebuild the project from scratch.
    Inspect the current implementation...'
    Guarantees it is completely eliminated from the generated docx.
    """
    out_file = str(tmp_path / "user_leakage_test.docx")
    cfg = TemplateConfig()
    
    leaked_input = (
        "# PROJECT UPDATE — EXPERIMENT HEADER + TYPOGRAPHY CORRECTION\n\n"
        "You are modifying the existing General College Laboratory Record Automation System.\n\n"
        "Do NOT rebuild the project from scratch.\n\n"
        "Inspect the current implementation and fix the document-generation/layout engine based on the actual reference template.\n\n"
        "from collections import OrderedDict\n\n"
        "def menu():\n"
        "    data = OrderedDict()\n"
        "    data['key'] = 'val'\n"
        "    print(data)\n"
    )
    
    exp = ExperimentData(
        experiment_number="3.A",
        title="Menu-Driven Program Using Ordered Dictionary",
        aim="To test prompt leakage prevention.",
        algorithm="1. Initialize dictionary.\n2. Add entries.\n3. Display.",
        coding=leaked_input,
        result="Success"
    )
    
    DocxGenerator.generate_new_record(exp, cfg, out_file)
    assert os.path.exists(out_file)
    
    doc = docx.Document(out_file)
    full_text = " ".join([p.text for p in doc.paragraphs])
    for tbl in doc.tables:
        full_text += " " + " ".join([c.text for r in tbl.rows for c in r.cells])
        
    assert "PROJECT UPDATE" not in full_text
    assert "General College Laboratory Record Automation System" not in full_text
    assert "Do NOT rebuild" not in full_text
    assert "Inspect the current implementation" not in full_text
    assert "from collections import OrderedDict" in full_text

def test_section_parser_sanitizes_pasted_prompt():
    """Verifies that SectionParser strips developer instructions from pasted notes."""
    from backend.parser.section_parser import SectionParser
    raw_pasted = (
        "# PROJECT UPDATE — EXPERIMENT HEADER + TYPOGRAPHY CORRECTION\n"
        "You are modifying the existing General College Laboratory Record Automation System.\n"
        "EX NO: 3.A\n"
        "TITLE: Menu-Driven Program Using Ordered Dictionary\n"
        "AIM:\n"
        "To perform insert, display, and search operations using an ordered dictionary in Python.\n"
        "ALGORITHM:\n"
        "1. Create an empty ordered dictionary.\n"
        "2. Display the menu.\n"
        "CODING:\n"
        "# PROJECT UPDATE — EXPERIMENT HEADER\n"
        "from collections import OrderedDict\n"
        "def menu():\n"
        "    data = OrderedDict()\n"
        "OUTPUT:\n"
        "1. Insert\n"
        "RESULT:\n"
        "Successfully executed.\n"
    )
    parsed = SectionParser.parse_raw_text(raw_pasted)
    assert parsed["experiment_number"] == "3.A"
    assert "PROJECT UPDATE" not in parsed["coding"]
    assert "from collections import OrderedDict" in parsed["coding"]
    assert "PROJECT UPDATE" not in parsed["aim"]

def test_prompt_only_input_uses_safe_academic_fallback():
    """Verifies that if an entire field is composed of prompt text, a safe fallback is applied."""
    prompt_only = (
        "# PROJECT UPDATE — EXPERIMENT HEADER + TYPOGRAPHY CORRECTION\n"
        "You are modifying the existing General College Laboratory Record Automation System.\n"
        "Do NOT rebuild the project from scratch.\n"
    )
    exp = ExperimentData(
        experiment_number="1.A",
        title=prompt_only,
        aim=prompt_only,
        algorithm=prompt_only,
        coding=prompt_only,
        result=prompt_only
    )
    assert exp.title == "LAB EXPERIMENT"
    assert "PROJECT UPDATE" not in exp.coding
    assert "PROJECT UPDATE" not in exp.aim
    assert "PROJECT UPDATE" not in exp.algorithm
    assert "PROJECT UPDATE" not in exp.result
    assert "executed successfully" in exp.coding.lower()
