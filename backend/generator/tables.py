"""
Table Manager: Clones or creates Master Evaluation Tables and Header Tables
with exact borders, widths, alignments, and typography.
Enforces fixed table layouts so external marks columns never shift or distort.
"""
import copy
import docx
from docx.table import Table
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from ..models.schemas import ExperimentData, TemplateConfig

def set_cell_width(cell, width_dxa: int):
    """Sets explicit cell width in dxa (1 inch = 1440 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is not None:
        tcPr.remove(tcW)
    tcW_elem = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW_elem)

def set_cell_v_align(cell, align: str = "center"):
    """Sets cell vertical alignment (top, center, bottom)."""
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = tcPr.find(qn('w:vAlign'))
    if vAlign is not None:
        tcPr.remove(vAlign)
    vAlign_elem = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{align}"/>')
    tcPr.append(vAlign_elem)

def apply_cell_border(cell, **kwargs):
    """Applies clean single borders to a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    existing_borders = tcPr.find(qn('w:tcBorders'))
    if existing_borders is not None:
        tcPr.remove(existing_borders)
    tcBorders = parse_xml(f"""
        <w:tcBorders {nsdecls('w')}>
            <w:top w:val="{kwargs.get('top', 'single')}" w:sz="{kwargs.get('sz', '6')}" w:space="0" w:color="{kwargs.get('color', '000000')}"/>
            <w:left w:val="{kwargs.get('left', 'single')}" w:sz="{kwargs.get('sz', '6')}" w:space="0" w:color="{kwargs.get('color', '000000')}"/>
            <w:bottom w:val="{kwargs.get('bottom', 'single')}" w:sz="{kwargs.get('sz', '6')}" w:space="0" w:color="{kwargs.get('color', '000000')}"/>
            <w:right w:val="{kwargs.get('right', 'single')}" w:sz="{kwargs.get('sz', '6')}" w:space="0" w:color="{kwargs.get('color', '000000')}"/>
        </w:tcBorders>
    """)
    tcPr.append(tcBorders)

from xml.sax.saxutils import escape as escape_xml

def insert_table_in_doc(doc: docx.Document, tbl_elem) -> Table:
    """Inserts a table element at the current document position (before sectPr)."""
    if len(doc.paragraphs) == 1 and not doc.paragraphs[0].text:
        p0 = doc.paragraphs[0]._p
        p0.addprevious(tbl_elem)
        p0.getparent().remove(p0)
    else:
        p_ph = doc.add_paragraph()
        p_ph._p.addprevious(tbl_elem)
        p_ph._p.getparent().remove(p_ph._p)
    return Table(tbl_elem, doc)

class TableManager:
    @staticmethod
    def create_header_table(doc: docx.Document, data: ExperimentData, config: TemplateConfig):
        """Creates the Master Header Table matching the reference template geometry exactly."""
        title_text = data.title.strip() if data.title else ""
        subtitle_text = data.subtitle.strip() if data.subtitle else ""
        has_subtitle = bool(subtitle_text and subtitle_text.upper() != title_text.upper())

        exp_no = data.experiment_number.strip() if data.experiment_number else ""
        date_str = data.date.strip() if data.date else ""

        if has_subtitle:
            title_xml = f"""
            <w:r>
              <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                <w:b/>
                <w:sz w:val="28"/>
              </w:rPr>
              <w:t>{escape_xml(title_text)}</w:t>
              <w:br/>
              <w:t>{escape_xml(subtitle_text)}</w:t>
            </w:r>
            """
        else:
            title_xml = f"""
            <w:r>
              <w:rPr>
                <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                <w:b/>
                <w:sz w:val="28"/>
              </w:rPr>
              <w:t>{escape_xml(title_text)}</w:t>
            </w:r>
            """

        header_tbl_xml = f"""
        <w:tbl {nsdecls('w')}>
          <w:tblPr>
            <w:tblW w:type="auto" w:w="0"/>
            <w:jc w:val="center"/>
            <w:tblLayout w:type="fixed"/>
            <w:tblLook w:firstColumn="1" w:firstRow="1" w:lastColumn="0" w:lastRow="0" w:noHBand="0" w:noVBand="1" w:val="04A0"/>
          </w:tblPr>
          <w:tblGrid>
            <w:gridCol w:w="2880"/>
            <w:gridCol w:w="6869"/>
          </w:tblGrid>
          <w:tr>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="2880" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60" w:line="240" w:lineRule="auto"/>
                </w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="22"/>
                  </w:rPr>
                  <w:t>EX NO:{escape_xml(exp_no)}</w:t>
                </w:r>
              </w:p>
            </w:tc>
            <w:tc>
              <w:tcPr>
                <w:vMerge w:val="restart"/>
                <w:tcW w:w="6869" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="80" w:after="80" w:line="276" w:lineRule="auto"/>
                  <w:jc w:val="center"/>
                </w:pPr>
                {title_xml}
              </w:p>
            </w:tc>
          </w:tr>
          <w:tr>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="2880" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60" w:line="240" w:lineRule="auto"/>
                </w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="22"/>
                  </w:rPr>
                  <w:t>DATE:{escape_xml(date_str)}</w:t>
                </w:r>
              </w:p>
            </w:tc>
            <w:tc>
              <w:tcPr>
                <w:vMerge/>
                <w:tcW w:w="6869" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p/>
            </w:tc>
          </w:tr>
        </w:tbl>
        """
        tbl_elem = parse_xml(header_tbl_xml)
        return insert_table_in_doc(doc, tbl_elem)

    @staticmethod
    def create_evaluation_table(doc: docx.Document, config: TemplateConfig):
        """Creates the Master Evaluation Marks Table matching Table 1/Table 5 of the reference template."""
        eval_tbl_xml = f"""
        <w:tbl {nsdecls('w')}>
          <w:tblPr>
            <w:tblW w:type="auto" w:w="0"/>
            <w:jc w:val="right"/>
            <w:tblLook w:firstColumn="1" w:firstRow="1" w:lastColumn="0" w:lastRow="0" w:noHBand="0" w:noVBand="1" w:val="04A0"/>
            <w:tblLayout w:type="fixed"/>
          </w:tblPr>
          <w:tblGrid>
            <w:gridCol w:w="3600"/>
            <w:gridCol w:w="2160"/>
          </w:tblGrid>
          <w:tr>
            <w:trPr>
              <w:trHeight w:val="400" w:hRule="atLeast"/>
            </w:trPr>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="3600" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60"/>
                </w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="20"/>
                  </w:rPr>
                  <w:t>PROGRAM AND EXECUTION</w:t>
                </w:r>
              </w:p>
            </w:tc>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="2160" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60"/>
                </w:pPr>
              </w:p>
            </w:tc>
          </w:tr>
          <w:tr>
            <w:trPr>
              <w:trHeight w:val="400" w:hRule="atLeast"/>
            </w:trPr>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="3600" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60"/>
                </w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="20"/>
                  </w:rPr>
                  <w:t>CLASS PERFORMANCE</w:t>
                </w:r>
              </w:p>
            </w:tc>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="2160" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60"/>
                </w:pPr>
              </w:p>
            </w:tc>
          </w:tr>
          <w:tr>
            <w:trPr>
              <w:trHeight w:val="400" w:hRule="atLeast"/>
            </w:trPr>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="3600" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60"/>
                </w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="20"/>
                  </w:rPr>
                  <w:t>VIVA</w:t>
                </w:r>
              </w:p>
            </w:tc>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="2160" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60"/>
                </w:pPr>
              </w:p>
            </w:tc>
          </w:tr>
          <w:tr>
            <w:trPr>
              <w:trHeight w:val="400" w:hRule="atLeast"/>
            </w:trPr>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="3600" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60"/>
                </w:pPr>
                <w:r>
                  <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="20"/>
                  </w:rPr>
                  <w:t>TOTAL</w:t>
                </w:r>
              </w:p>
            </w:tc>
            <w:tc>
              <w:tcPr>
                <w:tcW w:w="2160" w:type="dxa"/>
                <w:vAlign w:val="center"/>
                <w:tcBorders>
                  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                </w:tcBorders>
              </w:tcPr>
              <w:p>
                <w:pPr>
                  <w:spacing w:before="60" w:after="60"/>
                </w:pPr>
              </w:p>
            </w:tc>
          </w:tr>
        </w:tbl>
        """
        tbl_elem = parse_xml(eval_tbl_xml)
        return insert_table_in_doc(doc, tbl_elem)
